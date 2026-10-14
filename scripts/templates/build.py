"""
Builds the free templates (lib/site/templates.ts, D-37) into private/templates/.

The words follow the articles they go with: the charter follows "leadership-team-charter" and
"leadership-team-purpose", the agenda "leadership-team-meeting-agenda". The outputs are committed, like the
share cards, so a deploy needs no Python.

    python3 scripts/templates/build.py

Needs python-docx. The product's name is read from lib/brand.ts, so a new name is a rebuild, not an edit.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "private" / "templates"
PRODUCT = re.search(r"PRODUCT_NAME = '([^']+)'", (ROOT / "lib" / "brand.ts").read_text()).group(1)
SITE = "www.leaderpact.com"

INK = RGBColor(0x11, 0x11, 0x11)
TEXT2 = RGBColor(0x55, 0x52, 0x4A)
SHADE = "EFE9DF"  # --rule-soft, the design's chip and rule colour


def new_document(title: str, intro: str, landscape: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_height, section.page_width = Cm(21.0), Cm(29.7)
    else:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.0))
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)

    footer = section.footer.paragraphs[0]
    footer.text = f"{PRODUCT} · {SITE} · Free to use, change and share."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = TEXT2

    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(intro)
    for run in p.runs:
        run.font.color.rgb = TEXT2
    return doc


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def table(doc: Document, head: list[str], rows: list[list[str]], widths: list[float] | None = None, blank_rows: int = 0):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.text = h
        shade(cell, SHADE)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(9.5)
    for row in rows + [[""] * len(head)] * blank_rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9.5)
    if widths:
        # Word reads the cells' widths, LibreOffice the grid's: set both, and stop autofit undoing them
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def lines(doc: Document, prompt: str, count: int = 3) -> None:
    doc.add_paragraph(prompt, style="Heading 2")
    for _ in range(count):
        doc.add_paragraph("_" * (124 if doc.sections[0].orientation == WD_ORIENT.LANDSCAPE else 86)).runs[0].font.color.rgb = TEXT2


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def charter() -> None:
    doc = new_document(
        "Leadership team charter",
        "How to use it: each member answers sections 1 to 3 alone first. Then compare the answers in a working "
        "session and agree one version. Keep the charter to two pages, and review it at least twice a year.",
    )

    doc.add_heading("1. Our purpose", level=1)
    lines(doc, "Why does this team exist? What does it add that its members could not add alone?")
    doc.add_paragraph("Which purposes matter most for us now? Weigh each one: high, medium or low.", style="Heading 2")
    table(
        doc,
        ["Purpose", "Weight", "What it means for our agenda"],
        [
            ["Set direction for the organisation", "", ""],
            ["Make better decisions together", "", ""],
            ["Get decisions carried out", "", ""],
            ["Work as one across units and levels", "", ""],
            ["Lead change", "", ""],
            ["Develop capable leaders", "", ""],
        ],
        [7.0, 2.4, 7.6],
    )

    doc.add_heading("2. Our team tasks and decision rights", level=1)
    doc.add_paragraph("List the issues that need the whole team, and say who decides each one.")
    table(
        doc,
        ["Task or issue", "Why it needs the whole team", "We decide / we advise / others decide", "Owner"],
        [],
        [5.0, 5.2, 4.2, 2.6],
        blank_rows=6,
    )
    lines(doc, "Issues that do not belong here, and where they go instead (a one-to-one, a smaller group, a written update)", 2)

    doc.add_heading("3. Our members", level=1)
    table(doc, ["Core member", "Role", "Why they are in the core team"], [], [4.5, 4.5, 8.0], blank_rows=6)
    table(doc, ["Extended member or guest", "When they join", "For which tasks"], [], [4.5, 4.5, 8.0], blank_rows=3)

    doc.add_heading("4. How we work together", level=1)
    doc.add_paragraph("Agree four to six working agreements. For each, say how you will know whether you live by it.")
    table(
        doc,
        ["Working agreement", "How we will know"],
        [
            ["Example: we say what we think in the meeting, not after it.", "Example: no decision is reopened in the corridor."],
            ["Example: papers go out two working days before the meeting.", "Example: the chair moves an item whose papers are late."],
        ],
        [8.5, 8.5],
        blank_rows=4,
    )

    doc.add_heading("5. What we need to succeed", level=1)
    table(
        doc,
        ["Area", "What we have", "The gap, and what we will do"],
        [["Information", "", ""], ["Time", "", ""], ["Recognition and rewards", "", ""], ["Support and development", "", ""]],
        [4.0, 6.0, 7.0],
    )

    doc.add_heading("6. Our key stakeholders", level=1)
    table(doc, ["Who", "What they need from us", "Who owns the relationship"], [], [5.0, 7.5, 4.5], blank_rows=4)

    doc.add_heading("7. Review", level=1)
    table(doc, ["Agreed on", "Agreed by", "Next review on"], [["", "", ""]], [4.0, 9.0, 4.0])

    doc.save(OUT / "leadership-team-charter-template.docx")


def agenda() -> None:
    doc = new_document(
        "Leadership team meeting agenda",
        "How to use it: item owners fill in their rows and attach the pre-read by the deadline. The chair runs the "
        "check below and moves items that fail it. In the meeting, read each item's type and desired outcome aloud first.",
        landscape=True,
    )

    table(
        doc,
        ["Meeting", "Date and time", "Place or link", "Chair", "Notes by", "Papers by"],
        [["", "", "", "", "", ""]],
        [5.2, 4.2, 4.2, 4.0, 4.0, 4.1],
    )
    lines(doc, "The purpose of this meeting, in one sentence", 1)

    doc.add_heading("Agenda", level=1)
    table(
        doc,
        ["#", "Item", "Type", "Desired outcome (one sentence)", "Team task", "Owner", "Time", "Pre-read"],
        [
            ["1", "Example: follow-up of last meeting's actions", "Follow-up", "Each open action is on track, or has a new date.", "", "", "10 min", "Action list"],
            ["2", "Example: priorities for the next quarter", "Decision", "We agree the three priorities for the next quarter.", "", "", "40 min", "Options paper"],
            ["3", "Example: effects of a hiring pause", "Discussion", "We understand the consequences for each unit before we decide next month.", "", "", "30 min", "Unit notes"],
            ["4", "Example: customer survey results", "Information", "Everyone knows the main findings and has asked their questions.", "", "", "15 min", "Summary"],
        ],
        [0.8, 5.2, 2.3, 7.0, 3.0, 2.4, 1.8, 3.2],
        blank_rows=4,
    )

    doc.add_heading("Before an item goes on the agenda", level=1)
    bullets(
        doc,
        [
            "Does it need the whole team, or would a one-to-one, a smaller group or a written update do?",
            "Is its type clear: information, discussion, decision or follow-up?",
            "Can the owner write the desired outcome in one sentence?",
            "Will the pre-read be ready by the deadline?",
        ],
    )

    doc.add_heading("The four item types", level=1)
    table(
        doc,
        ["Type", "What the owner wants from the team"],
        [
            ["Information", "The team knows something, with time for questions. No decision."],
            ["Discussion", "Views and options on an issue. No decision today."],
            ["Decision", "A choice the team makes in this meeting, recorded with an owner."],
            ["Follow-up", "A check that earlier decisions are being carried out."],
        ],
        [5.0, 20.7],
    )

    doc.add_heading("Close of the meeting", level=1)
    table(doc, ["Decision", "Owner", "Date"], [], [17.7, 5.0, 3.0], blank_rows=4)
    table(doc, ["Action", "Owner", "Due"], [], [17.7, 5.0, 3.0], blank_rows=4)
    lines(doc, "Parked items, and where they go next", 2)
    doc.add_paragraph("Two minutes on the meeting itself", style="Heading 2")
    table(
        doc,
        ["Question", "Yes / partly / no", "What we will do differently"],
        [
            ["Did each item reach its desired outcome?", "", ""],
            ["Did the discussion stay on the item at hand?", "", ""],
            ["Did every item need the whole team?", "", ""],
        ],
        [10.0, 4.5, 11.2],
    )

    doc.save(OUT / "leadership-team-meeting-agenda-template.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    charter()
    agenda()
    for f in sorted(OUT.iterdir()):
        print(f"{f.relative_to(ROOT)}: {f.stat().st_size} bytes")
